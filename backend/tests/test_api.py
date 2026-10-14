"""End-to-end API tests (FastAPI TestClient on a throwaway SQLite database).

Covers the audit fixes: app startup (vendors router), role enforcement,
shipment workflow rules, catalogue upload/download, quotes access, mosquito
net pricing, project relinking and safe file downloads.

Run:  cd backend && python -m pytest tests -v
"""
import base64

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402  (seeds the temp DB)

client = TestClient(app)


def _login(email, pwd):
    r = client.post("/api/auth/login", json={"email": email, "password": pwd})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture(scope="module")
def admin():
    return _login("joshua.n@interdecng.com", "123456")


@pytest.fixture(scope="module")
def sales():
    return _login("sales@interdecng.com", "sales123")


@pytest.fixture(scope="module")
def viewer():
    return _login("viewer@interdecng.com", "viewer123")


def _b64(data: bytes) -> str:
    return base64.b64encode(data).decode()


def test_health_and_core_lists(admin):
    assert client.get("/api/health").json()["status"] == "ok"
    for path in ("/api/vendors", "/api/shippers", "/api/projects", "/api/shipments"):
        r = client.get(path, headers=admin)
        assert r.status_code == 200 and isinstance(r.json(), list) and r.json(), path


def test_unauthenticated_rejected():
    assert client.get("/api/vendors").status_code == 401
    assert client.get("/api/quotes").status_code == 401


def test_viewer_is_read_only(viewer):
    assert client.get("/api/shipments", headers=viewer).status_code == 200
    assert client.post("/api/projects", json={"name": "X"}, headers=viewer).status_code == 403
    assert client.delete("/api/shipments/sp1", headers=viewer).status_code == 403
    assert client.post("/api/vendors", json={"name": "X"}, headers=viewer).status_code == 403
    # viewer has neither quotes nor facade access
    assert client.get("/api/quotes", headers=viewer).status_code == 403
    # user list is admin-only
    assert client.get("/api/users", headers=viewer).status_code == 403


def test_vendor_crud_and_catalogue_roundtrip(sales):
    v = client.post("/api/vendors", json={"name": "  Test Glass Co  ", "category": "Glass & Glazing"}, headers=sales).json()
    assert v["name"] == "Test Glass Co" and v["catalogs"] == []
    pdf = b"%PDF-1.4 test"
    r = client.post(f"/api/vendors/{v['id']}/catalogs",
                    json=[{"name": "../evil\"name.pdf", "size": 1, "type": "application/pdf", "data": _b64(pdf)}],
                    headers=sales)
    assert r.status_code == 200, r.text
    cat = r.json()[0]
    assert "data" not in cat and cat["size"] == len(pdf) and "/" not in cat["name"] and '"' not in cat["name"]
    d = client.get(f"/api/vendors/{v['id']}/catalogs/0/download", headers=sales)
    assert d.status_code == 200 and d.content == pdf and d.headers["content-type"] == "application/pdf"
    assert client.delete(f"/api/vendors/{v['id']}/catalogs/0", headers=sales).json() == {"ok": True}
    assert client.get(f"/api/vendors/{v['id']}/catalogs", headers=sales).json() == []
    assert client.delete(f"/api/vendors/{v['id']}", headers=sales).status_code == 200


def test_vendor_in_use_cannot_be_deleted(admin):
    r = client.delete("/api/vendors/vn6", headers=admin)
    assert r.status_code == 400


def test_html_upload_is_never_served_inline(sales):
    v = client.post("/api/vendors", json={"name": "XSS Co"}, headers=sales).json()
    client.post(f"/api/vendors/{v['id']}/catalogs",
                json=[{"name": "x.html", "type": "text/html", "data": _b64(b"<script>alert(1)</script>")}],
                headers=sales)
    d = client.get(f"/api/vendors/{v['id']}/catalogs/0/download", headers=sales)
    assert d.headers["content-type"] == "application/octet-stream"
    assert d.headers["content-disposition"].startswith("attachment")


def test_invalid_base64_rejected(sales):
    v = client.post("/api/vendors", json={"name": "B64 Co"}, headers=sales).json()
    r = client.post(f"/api/vendors/{v['id']}/catalogs",
                    json=[{"name": "a.pdf", "type": "application/pdf", "data": "not base64!!"}], headers=sales)
    assert r.status_code == 400


def _file(name="doc.pdf"):
    return {"name": name, "size": 3, "type": "application/pdf", "data": _b64(b"abc")}


def test_shipment_workflow_is_enforced(sales):
    s = client.post("/api/shipments", json={
        "companyId": "facade", "category": "Hardware", "vendorId": "vn7", "description": "Workflow test",
    }, headers=sales).json()
    sid = s["id"]
    assert s["status"] == "order_placed" and s["ref"].startswith("IMP-")

    put = lambda body: client.put(f"/api/shipments/{sid}", json=body, headers=sales)  # noqa: E731
    assert put({"status": "in_transit"}).status_code == 400          # cannot skip production
    assert put({"packingList": _file()}).status_code == 400          # invoice first
    assert put({"vendorInvoice": _file()}).json()["status"] == "under_production"
    r = put({"shipperId": "sh1", "paymentTerms": "Before Shipment", "shipMethod": "FCL", "eta": "01/01/27"})
    assert r.json()["status"] == "under_production"                  # saving shipper no longer ships
    assert put({"status": "in_transit"}).status_code == 400          # needs packing list
    put({"packingList": _file()})
    assert put({"status": "in_transit"}).json()["status"] == "in_transit"
    assert put({"goodsConfirmed": True}).status_code == 400          # needs shipper invoice
    put({"shipperInvoice": _file()})
    assert "shipping cost" in put({"goodsConfirmed": True}).json()["detail"]      # costs come first
    costs = lambda body: client.put(f"/api/shipments/{sid}/costs", json=body, headers=sales)  # noqa: E731
    assert costs({"shipping": {"currency": "USD", "lines": [{"label": "Freight", "amount": 1200}]}}).json()["shippingCost"]["total"] == 1200
    assert "clearing" in put({"goodsConfirmed": True}).json()["detail"]
    assert costs({"clearing": {"lines": [{"label": "Duty", "amount": 50000}]}}).status_code == 400   # agent first
    c = costs({"clearingAgentId": "ca1", "clearing": {"currency": "NGN", "lines": [{"label": "Duty", "amount": 50000}, {"label": "Agent fee", "amount": 15000}], "paid": True, "paidOn": "2026-10-01"}}).json()
    assert c["clearingCost"]["total"] == 65000 and c["clearingAgentId"] == "ca1" and c["dates"].get("clearingCosted")
    done = put({"goodsConfirmed": True}).json()
    assert done["status"] == "completed" and done["goodsConfirmed"] and done["dates"].get("closed")
    assert put({"status": "order_placed"}).status_code == 400        # non-admin cannot reopen


def test_invalid_status_value_rejected(admin):
    assert client.put("/api/shipments/sp5", json={"status": "teleported"}, headers=admin).status_code == 422


def test_shipment_refs_unique_after_delete(sales):
    body = {"companyId": "doortec", "category": "Hardware", "vendorId": "vn7", "description": "Ref test"}
    a = client.post("/api/shipments", json=body, headers=sales).json()
    b = client.post("/api/shipments", json=body, headers=sales).json()
    client.delete(f"/api/shipments/{a['id']}", headers=sales)
    c = client.post("/api/shipments", json=body, headers=sales).json()
    refs = [s["ref"] for s in client.get("/api/shipments", headers=sales).json()]
    assert len(refs) == len(set(refs)) and c["ref"] != b["ref"]


def test_quotes_access_and_project_relink(sales):
    q = client.post("/api/quotes", json={"clientName": "Relink Ltd", "items": []}, headers=sales).json()
    r = client.put(f"/api/quotes/{q['id']}", json={"projectId": "pj2"}, headers=sales).json()
    assert r["projectId"] == "pj2" and r["projectName"] == "Tel Aviv Marina Hotel"
    assert client.put(f"/api/quotes/{q['id']}", json={"projectId": "nope"}, headers=sales).status_code == 400


def test_mosquito_net_rate_is_applied(admin):
    item = {"id": "m1", "product": "et_window", "subTypeId": "etw_s2", "width": 2000, "height": 1000, "qty": 1,
            "glassTypeId": "single", "glassThickness": "6mm", "glassColour": "clear", "hasMosq": True}
    assert client.post("/api/quotes/calculate", json={"items": [item]}, headers=admin).json()["items"][0]["bom"]["mosqCost"] == 0
    cat = client.get("/api/quotes/catalog", headers=admin).json()["values"]
    rates = dict(cat["mosqRates"], etw_s2=10000)
    assert client.put("/api/quotes/catalog/mosqRates", json={"value": rates}, headers=admin).status_code == 200
    bom = client.post("/api/quotes/calculate", json={"items": [item]}, headers=admin).json()["items"][0]["bom"]
    assert bom["mosqCost"] == 10000  # 2 m2 x 0.5 (sliding half-area) x 10,000


def test_saved_revision_keeps_its_priced_snapshot(admin):
    item = {"id": "g1", "product": "et_window", "subTypeId": "etw_single", "width": 1200, "height": 1500, "qty": 10,
            "glassTypeId": "tempered", "glassThickness": "8mm", "glassColour": "clear"}
    q = client.post("/api/quotes", json={"clientName": "Snapshot", "items": [item], "markupPct": 100, "vatPct": 7.5},
                    headers=admin).json()
    before = q["latest"]["items"][0]["bom"]["total"]
    cat = client.get("/api/quotes/catalog", headers=admin).json()["values"]
    r2 = dict(cat["rates"], fabRate=99999)
    client.put("/api/quotes/catalog/rates", json={"value": r2}, headers=admin)
    after = client.get(f"/api/quotes/{q['id']}", headers=admin).json()["latest"]["items"][0]["bom"]["total"]
    client.put("/api/quotes/catalog/rates", json={"value": cat["rates"]}, headers=admin)
    assert before == after == 1666000


def test_catalog_validation(admin):
    bad = {"defaultMarkup": 50, "minMarkup": 70, "vatRate": 7.5, "validity": 14}
    assert client.put("/api/quotes/catalog/settings", json={"value": bad}, headers=admin).status_code == 400
    assert client.put("/api/quotes/catalog/rates", json={"value": {"brandRates": {}}}, headers=admin).status_code == 400


def test_sent_revision_is_locked(admin):
    q = client.post("/api/quotes", json={"clientName": "Lock", "items": []}, headers=admin).json()
    rid = q["latest"]["id"]
    client.post(f"/api/quotes/{q['id']}/revisions/{rid}/status", json={"status": "sent"}, headers=admin)
    r = client.put(f"/api/quotes/{q['id']}/revisions/{rid}", json={"markupPct": 150}, headers=admin)
    assert r.status_code == 400
    assert client.post(f"/api/quotes/{q['id']}/revisions/{rid}/status", json={"status": "bogus"},
                       headers=admin).status_code == 422


def test_quote_export_filename_header(admin):
    q = client.get("/api/quotes", headers=admin).json()[0]
    r = client.get(f"/api/quotes/{q['id']}/export", headers=admin)
    assert r.status_code == 200 and "filename*=" in r.headers["content-disposition"]


def test_change_password_keeps_session_and_validates(sales):
    r = client.post("/api/auth/change-password", json={"currentPassword": "wrong", "newPassword": "abcdef"}, headers=sales)
    assert r.status_code == 401
    r = client.post("/api/auth/change-password", json={"currentPassword": "sales123", "newPassword": " pass word "},
                    headers=sales)
    assert r.status_code == 200
    _login("sales@interdecng.com", " pass word ")  # spaces are preserved, not stripped
    # demo passwords can't be set again
    r = client.post("/api/auth/change-password", json={"currentPassword": " pass word ", "newPassword": "sales123"},
                    headers=sales)
    assert r.status_code == 400
    me = client.get("/api/auth/me", headers=sales).json()
    assert me["mustChangePassword"] is False


def test_user_admin_rules(admin):
    r = client.post("/api/users", json={"name": "New Person", "email": "new@x.io", "password": "secret1",
                                        "apps": {"quotes": {"access": True, "role": "sales"}, "bogus": {"access": True}}},
                    headers=admin)
    assert r.status_code == 200, r.text
    u = r.json()
    assert u["apps"]["quotes"]["access"] is True and "bogus" not in u["apps"]
    assert client.post("/api/users", json={"name": "Bad", "email": "b@x.io", "password": "secret1",
                                           "apps": {"importflow": {"access": True, "role": "god"}}},
                       headers=admin).status_code == 400
    assert client.put("/api/users/u_admin", json={"platformRole": "user"}, headers=admin).status_code == 400
    assert client.delete(f"/api/users/{u['id']}", headers=admin).status_code == 200


def test_demo_password_forces_change():
    r = client.post("/api/auth/login", json={"email": "viewer@interdecng.com", "password": "viewer123"})
    assert r.json()["user"]["mustChangePassword"] is True


def test_admin_created_user_must_change_password(admin):
    u = client.post("/api/users", json={"name": "Temp", "email": "temp@x.io", "password": "temp123"}, headers=admin).json()
    assert u["mustChangePassword"] is True
    client.delete(f"/api/users/{u['id']}", headers=admin)


def test_composite_item_priced_with_elements(admin):
    base = {"id": "c1", "product": "et_window", "subTypeId": "etw_single", "width": 1000, "height": 1500, "qty": 2,
            "glassTypeId": "tempered", "glassThickness": "8mm", "glassColour": "clear", "hasSubframe": True}
    el = {"id": "e1", "product": "et_window", "subTypeId": "etw_fixed", "width": 600, "height": 1500,
          "glassTypeId": "tempered", "glassThickness": "8mm", "glassColour": "clear"}
    single = client.post("/api/quotes/calculate", json={"items": [base]}, headers=admin).json()["cost"]
    comp = client.post("/api/quotes/calculate", json={"items": [dict(base, elements=[el])]}, headers=admin).json()
    bom = comp["items"][0]["bom"]
    assert bom["isComposite"] and comp["cost"] > single
    # shared-edge sub-frame rule: 1.5 x (W + H) per frame x qty
    assert bom["sfPerimLm"] == round((1.5 * (1.0 + 1.5) + 1.5 * (0.6 + 1.5)) * 2, 3)


def test_client_export_hides_costs_and_internal_is_admin_only(admin, sales):
    from openpyxl import load_workbook
    import io
    q = client.get("/api/quotes", headers=admin).json()[0]
    wb = load_workbook(io.BytesIO(client.get(f"/api/quotes/{q['id']}/export", headers=sales).content))
    assert wb.sheetnames == ["Work Breakdown"]
    assert client.get(f"/api/quotes/{q['id']}/export?internal=true", headers=sales).status_code == 403
    wb2 = load_workbook(io.BytesIO(client.get(f"/api/quotes/{q['id']}/export?internal=true", headers=admin).content))
    assert "BOM Details" in wb2.sheetnames


def test_won_lost_needs_reason_and_records_it(admin):
    q = client.post("/api/quotes", json={"clientName": "Outcome Ltd", "items": []}, headers=admin).json()
    rid = q["latest"]["id"]
    url = f"/api/quotes/{q['id']}/revisions/{rid}/status"
    client.post(url, json={"status": "sent"}, headers=admin)
    assert client.post(url, json={"status": "won"}, headers=admin).status_code == 400
    r = client.post(url, json={"status": "won", "note": "PO-4471"}, headers=admin).json()
    assert r["status"] == "won" and r["statusNote"] == "PO-4471"
    # undo back to sent clears the note
    assert client.post(url, json={"status": "sent"}, headers=admin).json()["statusNote"] == ""


def test_fx_rates_shared_and_admin_only(admin, viewer):
    assert client.get("/api/settings/fx", headers=viewer).json()["rates"]["USD"] == 1
    assert client.put("/api/settings/fx", json={"rates": {"EUR": 1.1}}, headers=viewer).status_code == 403
    r = client.put("/api/settings/fx", json={"rates": {"EUR": 1.1, "USD": 5}}, headers=admin).json()
    assert r["rates"]["EUR"] == 1.1 and r["rates"]["USD"] == 1 and r["updatedBy"]
    assert client.put("/api/settings/fx", json={"rates": {"EUR": -1}}, headers=admin).status_code == 400


def test_activity_hides_sign_ins_by_default(admin):
    acts = client.get("/api/activity?limit=50", headers=admin).json()
    assert all(not a["action"].startswith("auth.") for a in acts)
    assert any(a["action"].startswith("auth.") for a in client.get("/api/activity?limit=50&include_auth=true", headers=admin).json())


def _item(**kw):
    base = {"id": "i1", "product": "et_window", "subTypeId": "etw_fixed", "width": 1000, "height": 1000, "qty": 2,
            "glassTypeId": "single", "glassThickness": "6mm", "glassColour": "clear"}
    return {**base, **kw}


def test_rate_card_additions_flow_into_pricing(admin):
    cat = client.get("/api/quotes/catalog", headers=admin).json()["values"]
    rates, gc, acp = cat["rates"], cat["gc"], cat["acpPanels"]
    calc = lambda items: client.post("/api/quotes/calculate", headers=admin,
                                     json={"items": items, "markupPct": 100, "vatPct": 0}).json()
    base_total = calc([_item()])["items"][0]["bom"]["total"]
    try:
        # new glass type + thickness + colour
        gc2 = {**gc, "types": gc["types"] + [{"id": "g_refl", "label": "Reflective", "isDGU": False, "custom": True}],
               "thicknesses": {**gc["thicknesses"], "g_refl": ["6mm"]},
               "basePrices": {**gc["basePrices"], "g_refl__6mm": 20000},
               "colours": gc["colours"] + [{"id": "c_bronze", "label": "Bronze", "premium": 1000, "custom": True}]}
        assert client.put("/api/quotes/catalog/gc", headers=admin, json={"value": gc2}).status_code == 200
        bom = calc([_item(glassTypeId="g_refl", glassColour="c_bronze")])["items"][0]["bom"]
        assert bom["glassCost"] == round(2.0 * 21000)

        # new ACP panel
        acp2 = acp + [{"id": "ap_gold", "name": "4mm Gold", "p": 30000, "custom": True}]
        assert client.put("/api/quotes/catalog/acpPanels", headers=admin, json={"value": acp2}).status_code == 200
        bom = calc([_item(product="acp", subTypeId="acp_g", acpPanelId="ap_gold")])["items"][0]["bom"]
        assert bom["glassCost"] == 60000

        # extra charges: always-on per m2, optional per item, product-limited
        extras = [
            {"id": "x_scaf", "name": "Scaffolding", "rate": 1000, "basis": "m2", "products": [], "optional": False},
            {"id": "x_coat", "name": "Coating", "rate": 5000, "basis": "item", "products": [], "optional": True},
            {"id": "x_door", "name": "Door only", "rate": 999, "basis": "item", "products": ["et_door"], "optional": False},
        ]
        assert client.put("/api/quotes/catalog/rates", headers=admin, json={"value": {**rates, "extras": extras}}).status_code == 200
        bom = calc([_item()])["items"][0]["bom"]
        assert bom["extrasCost"] == 2000 and bom["total"] == base_total + 2000
        assert [x["name"] for x in bom["extras"]] == ["Scaffolding"]
        bom = calc([_item(extraIds=["x_coat"])])["items"][0]["bom"]
        assert bom["extrasCost"] == 2000 + 10000

        # composite: per-item charge counted once, area charge covers every element
        comp = _item(extraIds=["x_coat"], elements=[{"id": "e1", "subTypeId": "etw_fixed", "width": 1000, "height": 1000,
                                                     "glassTypeId": "single", "glassThickness": "6mm"}])
        bom = calc([comp])["items"][0]["bom"]
        assert bom["extrasCost"] == 4000 + 10000

        # bad input is rejected
        bad = [{"id": "x1", "name": "", "rate": 1, "basis": "m2"}]
        assert client.put("/api/quotes/catalog/rates", headers=admin, json={"value": {**rates, "extras": bad}}).status_code == 400
        bad = [{"id": "x1", "name": "A", "rate": 1, "basis": "weekly"}]
        assert client.put("/api/quotes/catalog/rates", headers=admin, json={"value": {**rates, "extras": bad}}).status_code == 400
        gc_bad = {**gc2, "thicknesses": {**gc2["thicknesses"], "g_refl": ["6mm", "8mm"]}}
        assert client.put("/api/quotes/catalog/gc", headers=admin, json={"value": gc_bad}).status_code == 400
        assert client.put("/api/quotes/catalog/acpPanels", headers=admin,
                          json={"value": [{"id": "a", "name": "x", "p": -1}]}).status_code == 400
    finally:
        client.put("/api/quotes/catalog/gc", headers=admin, json={"value": gc})
        client.put("/api/quotes/catalog/acpPanels", headers=admin, json={"value": acp})
        client.put("/api/quotes/catalog/rates", headers=admin, json={"value": rates})


def test_rate_card_is_admin_only(sales):
    r = client.put("/api/quotes/catalog/acpPanels", headers=sales, json={"value": []})
    assert r.status_code == 403


def test_currencies_on_rate_card_and_quotes(admin):
    cat = client.get("/api/quotes/catalog", headers=admin).json()
    ngn_per_usd = cat["fx"]["USD"]
    assert cat["fx"]["NGN"] == 1 and ngn_per_usd > 100
    calc = lambda **kw: client.post("/api/quotes/calculate", headers=admin,
                                    json={"items": [_item()], "markupPct": 100, "vatPct": 0, **kw}).json()
    ngn = calc()
    assert ngn["currency"] == "NGN" and ngn["fxRate"] == 1

    # quote in USD at a fixed rate: every figure is the naira figure / rate (± whole-unit rounding)
    usd = calc(currency="USD", fxRate=1500)
    assert usd["currency"] == "USD" and usd["fxRate"] == 1500
    assert abs(usd["cost"] - ngn["cost"] / 1500) <= len(_item()) and usd["cost"] < ngn["cost"]
    assert client.post("/api/quotes/calculate", headers=admin, json={"items": [], "currency": "XYZ"}).status_code == 400

    # a rate-card price entered in USD is converted to naira for naira quotes
    try:
        r = client.put("/api/quotes/catalog/currencies", headers=admin,
                       json={"value": {"gc.basePrices.single__6mm": "USD"}})
        assert r.status_code == 200
        bom = calc()["items"][0]["bom"]
        assert abs(bom["glassCost"] - 2 * 8500 * ngn_per_usd) < 2
        # ...and stays 8,500 per m2 in a USD quote at today's rate
        bom = calc(currency="USD")["items"][0]["bom"]
        assert bom["glassCost"] == 17000
    finally:
        client.put("/api/quotes/catalog/currencies", headers=admin, json={"value": {}})

    # currency is stored on the quote and carried into new revisions
    q = client.post("/api/quotes", headers=admin, json={"clientName": "FX Ltd", "items": [_item()],
                                                         "currency": "USD", "fxRate": 1500}).json()
    assert q["currency"] == "USD" and q["latest"]["fxRate"] == 1500
    assert abs(q["totalNgn"] - q["total"] * 1500) < 1
    rev = client.post(f"/api/quotes/{q['id']}/revisions", headers=admin, json={}).json()
    assert rev["currency"] == "USD" and rev["fxRate"] == 1500
    upd = client.put(f"/api/quotes/{q['id']}/revisions/{rev['id']}", headers=admin, json={"currency": "NGN"}).json()
    assert upd["currency"] == "NGN" and upd["fxRate"] == 1 and upd["total"] > rev["total"] * 100
    x = client.get(f"/api/quotes/{q['id']}/export?rev_id={q['latest']['id']}", headers=admin)
    assert x.status_code == 200


def test_pdf_export_and_design(admin, sales):
    q = client.post("/api/quotes", headers=admin, json={"clientName": "PDF & Co <b>", "items": [_item(position="W1")]}).json()
    r = client.get(f"/api/quotes/{q['id']}/pdf", headers=sales)
    assert r.status_code == 200 and r.content[:5] == b"%PDF-"
    assert "attachment" in r.headers["content-disposition"] and q["number"] in r.headers["content-disposition"]
    assert "inline" in client.get(f"/api/quotes/{q['id']}/pdf?inline=true", headers=sales).headers["content-disposition"]

    from app import pdf_export
    rev = q["latest"]
    html = pdf_export.render_html(q, rev, {}, {"companyName": "Acme"})
    assert "PDF &amp; Co &lt;b&gt;" in html and "<b>" not in html.split("PDF")[1][:20]   # client data escaped
    assert "Acme" in html and str(int(rev["items"][0]["bom"]["total"])) not in html          # no cost figures

    # design: saved, validated, previewed (admin only)
    good = {"companyName": "Acme Glass", "brandColor": "#0A7E5C", "layout": "banner", "font": "serif",
            "show": {"sketch": True, "unitPrice": False}, "logo": "data:image/png;base64,iVBORw0KGgo="}
    assert client.put("/api/quotes/catalog/pdfDesign", headers=admin, json={"value": good}).status_code == 200
    try:
        p = client.post("/api/quotes/pdf/preview", headers=admin, json={"design": good})
        assert p.status_code == 200 and p.content[:5] == b"%PDF-"
        assert client.post("/api/quotes/pdf/preview", headers=sales, json={"design": good}).status_code == 403
        for bad in ({"brandColor": "red"}, {"layout": "fancy"}, {"logo": "data:text/html;base64,PGI+"},
                    {"logo": "https://evil.example/x.png"}, {"logoHeight": 999}):
            assert client.put("/api/quotes/catalog/pdfDesign", headers=admin, json={"value": bad}).status_code == 400, bad
    finally:
        client.put("/api/quotes/catalog/pdfDesign", headers=admin, json={"value": {}})
    d = client.get("/api/quotes/pdf/defaults", headers=sales).json()
    assert "classic" in d["layouts"] and d["design"]["title"] == "QUOTATION"


def test_importflow_exports_every_kind(admin, viewer):
    import openpyxl, io as _io
    ships = client.get("/api/shipments", headers=admin).json()
    projs = client.get("/api/projects", headers=admin).json()
    kinds = [("dashboard", ""), ("shipments", "&status=in_transit"), ("shipments", "&q=zzz-no-match"),
             ("shipment", f"&id={ships[0]['id']}"), ("projects", ""), ("project", f"&id={projs[0]['id']}"),
             ("vendors", ""), ("shippers", "")] + [("report", f"&type={t}") for t in ("status", "project", "company", "vendor", "shipper", "all")]
    for kind, extra in kinds:
        for fmt in ("pdf", "xlsx"):
            r = client.get(f"/api/importflow/export/{kind}?format={fmt}{extra}", headers=viewer)   # read-only users can export
            assert r.status_code == 200, (kind, fmt, r.text[:200])
            assert ".%s" % fmt in r.headers["content-disposition"]
            if fmt == "pdf":
                assert r.content[:5] == b"%PDF-"
            else:
                wb = openpyxl.load_workbook(_io.BytesIO(r.content))
                assert wb.active.max_row > 3
    # numbers stay numeric in Excel
    wb = openpyxl.load_workbook(_io.BytesIO(client.get("/api/importflow/export/shipments?format=xlsx", headers=admin).content))
    assert any(isinstance(c.value, (int, float)) and "#,##0" in c.number_format for row in wb.active.iter_rows() for c in row)
    assert client.get("/api/importflow/export/nope?format=pdf", headers=admin).status_code == 404
    assert client.get("/api/importflow/export/shipments?format=doc", headers=admin).status_code == 400
    assert client.get("/api/importflow/export/shipment?format=pdf&id=missing", headers=admin).status_code == 404


def test_importflow_document_design(admin, viewer):
    d = client.get("/api/importflow/design", headers=viewer).json()
    assert d["design"]["orientation"] == "auto"
    good = {"companyName": "Interdec Logistics", "brandColor": "#7C2D12", "layout": "banner", "orientation": "landscape",
            "show": {"stripes": False}}
    assert client.put("/api/importflow/design", headers=viewer, json={"design": good}).status_code == 403
    try:
        r = client.put("/api/importflow/design", headers=admin, json={"design": good})
        assert r.status_code == 200 and r.json()["design"]["show"]["summary"] is True
        p = client.post("/api/importflow/design/preview", headers=admin, json={"design": good})
        assert p.status_code == 200 and p.content[:5] == b"%PDF-" and "inline" in p.headers["content-disposition"]
        from app import doc_export
        html = doc_export.render_html({"title": "T <x>", "sections": []}, good)
        assert "Interdec Logistics" in html and "T &lt;x&gt;" in html and "landscape" in html
        for bad in ({"orientation": "sideways"}, {"brandColor": "blue"}, {"font": "comic"}):
            assert client.put("/api/importflow/design", headers=admin, json={"design": bad}).status_code == 400
    finally:
        client.put("/api/importflow/design", headers=admin, json={"design": {}})


class _FakeSMTP:
    sent = []

    def __init__(self, *a, **kw): pass
    def __enter__(self): return self
    def __exit__(self, *a): return False
    def ehlo(self): pass
    def has_extn(self, n): return False
    def login(self, u, p): pass
    def send_message(self, msg, to_addrs=None): _FakeSMTP.sent.append((msg["Subject"], list(to_addrs)))


def test_eta_alert_rules():
    from datetime import date
    from app import eta_alerts as E, models as M
    s = M.Shipment(ref="X-1", status="in_transit", eta="20/04/26",
                   eta_alerts={"enabled": True, "recipients": ["a@b.co"], "before": [7, 3, 1], "onDay": True, "overdue": True})
    assert E.parse_eta("20/04/26") == date(2026, 4, 20) == E.parse_eta("2026-04-20")
    assert E.due(s, date(2026, 4, 1)) == []                           # too early
    assert E.due(s, date(2026, 4, 13)) == [("before_7", 7)]
    assert E.due(s, date(2026, 4, 18)) == [("before_3", 2)]            # server was down: closest window only
    E.mark_sent(s, "before_3")
    assert E.due(s, date(2026, 4, 18)) == [] and E.due(s, date(2026, 4, 16)) == []   # 7-day is now stale too
    assert E.due(s, date(2026, 4, 19)) == [("before_1", 1)]
    assert E.due(s, date(2026, 4, 20)) == [("on_day", 0)]
    assert E.due(s, date(2026, 4, 23)) == [("overdue", -3)]
    s.eta = "2026-05-01"                                                 # new ETA re-arms reminders
    assert E.due(s, date(2026, 4, 28)) == [("before_3", 3)]
    s.status = "completed"
    assert E.due(s, date(2026, 5, 1)) == []
    for bad in ({"recipients": ["not-an-email"]}, {"before": [4]}, {"enabled": True, "recipients": []}):
        try:
            E.clean_settings(bad, dict(E.DEFAULTS))
            assert False, bad
        except ValueError:
            pass


def test_eta_alert_emails_end_to_end(admin, viewer, monkeypatch):
    import smtplib
    from datetime import date, timedelta
    sid = client.get("/api/shipments", headers=admin).json()[0]["id"]

    # email not set up: settings still save, test email explains why it can't send
    r = client.put(f"/api/shipments/{sid}/eta-alerts", headers=admin, json={"enabled": True, "recipients": ["ops@example.com"]})
    assert r.status_code == 200 and r.json()["email"]["configured"] is False
    t = client.post(f"/api/shipments/{sid}/eta-alerts/test", headers=admin, json={})
    assert t.status_code == 503 and "set up" in t.json()["detail"]
    assert client.put(f"/api/shipments/{sid}/eta-alerts", headers=viewer, json={"enabled": False}).status_code == 403
    assert client.put(f"/api/shipments/{sid}/eta-alerts", headers=admin, json={"recipients": ["bad@"]}).status_code == 400

    monkeypatch.setenv("INTERDEC_SMTP_HOST", "smtp.test")
    monkeypatch.setenv("INTERDEC_SMTP_FROM", "alerts@interdec.test")
    monkeypatch.setattr(smtplib, "SMTP", _FakeSMTP)
    _FakeSMTP.sent.clear()

    assert client.post(f"/api/shipments/{sid}/eta-alerts/test", headers=admin, json={}).status_code == 200
    assert _FakeSMTP.sent[-1][0].startswith("Test:") and _FakeSMTP.sent[-1][1] == ["ops@example.com"]

    # ETA in 2 days: the 3-day reminder goes once, never twice
    from app import eta_alerts as E
    eta = (E.today() + timedelta(days=2)).isoformat()
    client.put(f"/api/shipments/{sid}", headers=admin, json={"eta": eta})
    _FakeSMTP.sent.clear()
    first = client.post("/api/settings/eta-alerts/run", headers=admin).json()
    again = client.post("/api/settings/eta-alerts/run", headers=admin).json()
    subjects = [s for s, _ in _FakeSMTP.sent]
    assert first["sent"] >= 1 and any("arrives in 2 days" in s for s in subjects) and again["sent"] == 0

    # changing the ETA tells the list immediately
    _FakeSMTP.sent.clear()
    client.put(f"/api/shipments/{sid}", headers=admin, json={"eta": (E.today() + timedelta(days=20)).isoformat()})
    assert any(s.startswith("ETA changed for") for s, _ in _FakeSMTP.sent)

    info = client.get(f"/api/shipments/{sid}/eta-alerts", headers=viewer).json()
    assert info["enabled"] and info["log"][0]["status"] == "sent"
    assert client.get("/api/settings/email/log", headers=admin).status_code == 200
    assert client.get("/api/settings/email/log", headers=viewer).status_code == 403
    client.put(f"/api/shipments/{sid}/eta-alerts", headers=admin, json={"enabled": False})


def test_costs_rules_and_clearing_docs(admin, sales, viewer):
    sp2 = client.get("/api/shipments", headers=sales).json()
    sp2 = next(s for s in sp2 if s["id"] == "sp2")                       # seeded: in transit, shipping cost recorded
    assert sp2["shippingCost"]["total"] == 1525 and sp2["clearingAgentId"] == "ca2"
    url = "/api/shipments/sp2/costs"
    bad = [{"shipping": {"currency": "XYZ", "lines": []}}, {"shipping": {"lines": [{"label": "x", "amount": -5}]}},
           {"shipping": {"lines": [{"label": "", "amount": 10}]}}, {"clearing": {"paid": True, "paidOn": "yesterday"}},
           {"clearingAgentId": "nope"}]
    for b in bad:
        assert client.put(url, json=b, headers=sales).status_code == 400, b
    assert client.put(url, json={"shipping": {"lines": []}}, headers=viewer).status_code == 403
    assert client.put("/api/shipments/sp5/costs", json={"shipping": {"lines": []}}, headers=sales).status_code == 400  # not shipped yet

    # clearing documents: upload, download, remove
    r = client.post("/api/shipments/sp2/clearing-docs", json={"file": _file("paar.pdf"), "label": "PAAR"}, headers=sales)
    assert r.status_code == 200 and r.json()["clearingCost"]["docs"][0]["name"] == "paar.pdf"
    assert "data" not in r.json()["clearingCost"]["docs"][0]
    assert client.get("/api/shipments/sp2/clearing-docs/0", headers=viewer).status_code == 200
    assert client.get("/api/shipments/sp2/clearing-docs/5", headers=viewer).status_code == 404
    assert client.delete("/api/shipments/sp2/clearing-docs/0", headers=sales).json()["clearingCost"]["docs"] == []

    # clearing agents CRUD; an agent in use can't be deleted
    a = client.post("/api/clearing-agents", json={"name": "Tin Can Brokers", "country": "Nigeria"}, headers=sales).json()
    assert any(x["id"] == a["id"] for x in client.get("/api/clearing-agents", headers=viewer).json())
    assert client.delete("/api/clearing-agents/ca2", headers=sales).status_code == 400
    assert client.delete(f"/api/clearing-agents/{a['id']}", headers=sales).json() == {"ok": True}

    # exports include the costs
    for fmt in ("pdf", "xlsx"):
        assert client.get(f"/api/importflow/export/shipment?format={fmt}&id=sp1", headers=viewer).status_code == 200
        assert client.get(f"/api/importflow/export/report?format={fmt}&type=costs", headers=viewer).status_code == 200
        assert client.get(f"/api/importflow/export/agents?format={fmt}", headers=viewer).status_code == 200


def test_won_quote_creates_or_links_project(admin, sales):
    q = client.post("/api/quotes", headers=sales, json={"clientName": "Eko Atlantic Dev", "items": [_item()]}).json()
    rid = q["latest"]["id"]
    st = lambda body: client.post(f"/api/quotes/{q['id']}/revisions/{rid}/status", headers=sales, json=body)  # noqa: E731
    assert st({"status": "won", "note": "PO-1", "project": {"mode": "teleport"}}).status_code == 400
    assert st({"status": "won", "note": "PO-1", "project": {"mode": "link", "projectId": "missing"}}).status_code == 400
    r = st({"status": "won", "note": "PO-1", "project": {"mode": "create", "name": "Eko Atlantic Tower A"}})
    assert r.status_code == 200
    full = client.get(f"/api/quotes/{q['id']}", headers=sales).json()
    assert full["projectName"] == "Eko Atlantic Tower A"
    proj = next(p for p in client.get("/api/projects", headers=sales).json() if p["id"] == full["projectId"])
    assert proj["client"] == "Eko Atlantic Dev" and proj["status"] == "Active"
    b = client.get(f"/api/projects/{proj['id']}/budget", headers=sales).json()
    assert b["contract"] > 0 and b["wonQuotes"][0]["number"] == q["number"] and b["spent"] == 0

    # link mode
    q2 = client.post("/api/quotes", headers=sales, json={"clientName": "Second", "items": [_item()]}).json()
    client.post(f"/api/quotes/{q2['id']}/revisions/{q2['latest']['id']}/status", headers=sales,
                json={"status": "won", "note": "PO-2", "project": {"mode": "link", "projectId": proj["id"]}})
    b2 = client.get(f"/api/projects/{proj['id']}/budget", headers=sales).json()
    assert len(b2["wonQuotes"]) == 2 and b2["contract"] > b["contract"]


def test_budget_vs_actual_and_vendor_payments(admin, sales, viewer):
    b = client.get("/api/projects/budgets", headers=viewer).json()["pj1"]
    assert b["contract"] > 0 and b["spent"] > 0 and abs(b["margin"] - (b["contract"] - b["spent"])) < 1
    assert b["vendorPaid"] > 0 and b["goods"] + b["shipping"] + b["clearing"] - b["spent"] < 1

    sid = "sp4"   # under production, USD 12,400, nothing paid
    url = f"/api/shipments/{sid}/vendor-payments"
    assert client.post(url, headers=viewer, json={"amount": 100}).status_code == 403
    assert client.post(url, headers=sales, json={"amount": 0}).status_code == 400
    assert client.post(url, headers=sales, json={"amount": 99999}).status_code == 400          # more than owed
    assert client.post(url, headers=sales, json={"amount": 10, "date": "soon"}).status_code == 400
    s = client.post(url, headers=sales, json={"amount": 3720, "kind": "Deposit", "ref": "TT-1", "date": "2026-10-01"}).json()
    assert s["vendorPaid"] == 3720 and s["vendorBalance"] == 8680 and s["vendorPayStatus"] == "part"
    s = client.post(url, headers=sales, json={"amount": 8680, "kind": "Balance"}).json()
    assert s["vendorPayStatus"] == "paid" and s["vendorBalance"] == 0
    s = client.delete(f"{url}/{s['vendorPayments'][1]['id']}", headers=sales).json()
    assert s["vendorPayStatus"] == "part"
    assert client.delete(f"{url}/nope", headers=sales).status_code == 404


def test_budget_and_payables_exports(viewer):
    import openpyxl, io as _io
    for t in ("budget", "payables", "all"):
        for fmt in ("pdf", "xlsx"):
            r = client.get(f"/api/importflow/export/report?format={fmt}&type={t}", headers=viewer)
            assert r.status_code == 200, (t, fmt, r.text[:200])
    wb = openpyxl.load_workbook(_io.BytesIO(client.get("/api/importflow/export/report?format=xlsx&type=budget", headers=viewer).content))
    vals = [c.value for row in wb.active.iter_rows() for c in row if c.value is not None]
    assert "Azrieli Tower Renovation" in vals and "Contract" in vals
    for fmt in ("pdf", "xlsx"):
        assert client.get(f"/api/importflow/export/project?format={fmt}&id=pj1", headers=viewer).status_code == 200
        assert client.get(f"/api/importflow/export/projects?format={fmt}", headers=viewer).status_code == 200


def test_email_notifications(admin, sales, monkeypatch):
    import smtplib
    from datetime import datetime
    from app import notify_mail, models as M
    from app.database import SessionLocal
    me = client.get("/api/auth/me/notifications", headers=admin).json()
    assert me["prefs"]["payablesDigest"] is True and me["prefs"]["statusChanges"] is False
    r = client.put("/api/auth/me/notifications", headers=admin, json={"statusChanges": True, "quoteOutcomes": True, "bogus": 1})
    assert r.json()["prefs"]["statusChanges"] is True

    monkeypatch.setenv("INTERDEC_SMTP_HOST", "smtp.test")
    monkeypatch.setenv("INTERDEC_SMTP_FROM", "alerts@interdec.test")
    monkeypatch.setattr(smtplib, "SMTP", _FakeSMTP)
    _FakeSMTP.sent.clear()

    # status change by sales -> admin (opted in) is emailed; the person who did it is not
    s = client.post("/api/shipments", headers=sales, json={"companyId": "facade", "category": "Hardware", "vendorId": "vn7", "description": "Notify test", "value": 100}).json()
    client.put(f"/api/shipments/{s['id']}", headers=sales, json={"vendorInvoice": _file()})
    hit = [x for x in _FakeSMTP.sent if "is now Under Production" in x[0]]
    assert hit and "joshua.n@interdecng.com" in hit[0][1]

    # quote won -> outcome email
    q = client.post("/api/quotes", headers=sales, json={"clientName": "Mail Co", "items": [_item()]}).json()
    client.post(f"/api/quotes/{q['id']}/revisions/{q['latest']['id']}/status", headers=sales, json={"status": "won", "note": "PO-9"})
    assert any(x[0].startswith(f"{q['number']} won") for x in _FakeSMTP.sent)

    # digest: on demand, and the weekly one goes once per week, Monday morning only
    assert client.post("/api/auth/me/notifications/digest", headers=admin).status_code == 200
    assert any(x[0].startswith("Payables this week") for x in _FakeSMTP.sent)
    db = SessionLocal()
    try:
        db.query(M.FacadeSetting).filter(M.FacadeSetting.key == notify_mail.DIGEST_KEY).delete()
        db.commit()
        assert notify_mail.maybe_send_weekly_digest(db, datetime(2026, 10, 13, 9)) is False     # Tuesday
        assert notify_mail.maybe_send_weekly_digest(db, datetime(2026, 10, 12, 7)) is False     # Monday, too early
        assert notify_mail.maybe_send_weekly_digest(db, datetime(2026, 10, 12, 9)) is True
        assert notify_mail.maybe_send_weekly_digest(db, datetime(2026, 10, 12, 15)) is False    # already sent this week
    finally:
        db.close()
    client.put("/api/auth/me/notifications", headers=admin, json={"statusChanges": False})


def test_products_and_types_are_editable(admin):
    cat = client.get("/api/quotes/catalog", headers=admin).json()["values"]
    prods, subs, sf, kits = cat["products"], cat["subs"], cat["sf"], cat["hwKits"]
    put = lambda k, v: client.put(f"/api/quotes/catalog/{k}", headers=admin, json={"value": v})  # noqa: E731
    try:
        # new system with a brand of its own and a tilt & turn type
        new_p = prods + [{"id": "p_schuco", "label": "Schüco AWS 70 Window", "unit": "m2", "isLm": False, "brand": "SCHUCO", "custom": True}]
        assert put("products", new_p).status_code == 200
        new_s = {**subs, "p_schuco": [{"id": "t_tilt", "label": "Tilt & Turn", "kg": 13, "mosq": True, "subframe": True, "sliding": False, "custom": True}]}
        assert put("subs", new_s).status_code == 200
        assert put("sf", {**sf, "t_tilt": 13}).status_code == 200
        assert put("hwKits", {**kits, "t_tilt": [{"id": "hw1", "q": 1}, {"id": "hw2", "q": 2}]}).status_code == 200
        rates = cat["rates"]
        assert put("rates", {**rates, "brandRates": {**rates["brandRates"], "SCHUCO": 5000}}).status_code == 200
        it = {"id": "n1", "product": "p_schuco", "subTypeId": "t_tilt", "width": 1000, "height": 1000, "qty": 1,
              "glassTypeId": "single", "glassThickness": "6mm", "glassColour": "clear", "hasMosq": True, "hasSubframe": True}
        r = client.post("/api/quotes/calculate", headers=admin, json={"items": [it], "markupPct": 100, "vatPct": 0}).json()
        bom = r["items"][0]["bom"]
        assert bom["profileCost"] == 13 * 5000 and bom["hwCost"] == 5000 + 2 * 7000
        assert bom["subframeCost"] > 0 and bom["sealantCost"] > 0
        assert r["items"][0]["desc"] == "SCHUCO Tilt & Turn"
        # bad edits
        assert put("products", [p for p in prods if p["id"] != "acp"]).status_code == 400          # built-in ACP stays
        assert put("subs", {**new_s, "p_schuco": []}).status_code == 400
        assert put("subs", {**new_s, "p_schuco": [{"id": "etw_fixed", "label": "dup", "kg": 1}]}).status_code == 400
        assert put("hwKits", {"t_tilt": [{"id": "hw1", "q": -1}]}).status_code == 400
    finally:
        put("subs", subs); put("products", prods); put("sf", sf); put("hwKits", kits); put("rates", cat["rates"])


def _xlsx(rows):
    import openpyxl, io as _io, base64 as _b
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    buf = _io.BytesIO(); wb.save(buf)
    return {"name": "data.xlsx", "data": _b.b64encode(buf.getvalue()).decode()}


def test_import_vendors_projects_shipments_and_prices(admin, sales, viewer):
    import base64 as _b, openpyxl, io as _io
    t = client.get("/api/import/shipments/template", headers=sales)
    assert t.status_code == 200 and openpyxl.load_workbook(_io.BytesIO(t.content)).active.cell(1, 2).value.startswith("Company")
    assert client.get("/api/import/glass/template", headers=sales).status_code == 403        # price lists: admins
    assert client.post("/api/import/vendors/preview", headers=viewer, json={"file": _xlsx([["Name"]])}).status_code == 403

    # vendors: one new, one update (existing name, different case), one bad (no name)
    f = _xlsx([["Company Name", "Country", "E-mail"], ["Lagos Hardware Co", "Nigeria", "x@lh.ng"],
               ["vetro italia glass", "Italy", "new@vetro.it"], ["", "Ghana", ""]])
    p = client.post("/api/import/vendors/preview", headers=sales, json={"file": f}).json()
    assert p["summary"] == {"total": 3, "create": 1, "update": 1, "errors": 1}
    assert client.post("/api/import/vendors/commit", headers=sales, json={"file": f}).json() == {"created": 1, "updated": 1, "skipped": 1}
    vs = client.get("/api/vendors", headers=sales).json()
    assert any(v["name"] == "Lagos Hardware Co" for v in vs) and next(v for v in vs if v["name"] == "Vetro Italia Glass")["email"] == "new@vetro.it"

    # CSV works too, with ; separator
    csvf = {"name": "p.csv", "data": _b.b64encode("Project name;Client;Status\nIkoyi Villas;Ikoyi Dev;Active\n".encode()).decode()}
    assert client.post("/api/import/projects/commit", headers=sales, json={"file": csvf}).json()["created"] == 1

    # shipments: references existing vendor/project/shipper, records a payment
    f = _xlsx([["Company", "Vendor", "Description", "Project", "Value", "Currency", "Status", "ETA", "Shipper", "Paid to vendor", "Payment terms"],
               ["Facade", "Lagos Hardware Co", "Hinges batch", "Ikoyi Villas", 5000, "USD", "In Transit", "20/11/2026", "Global Freight Logistics", 2000, "After Shipment"],
               ["Facade", "Nobody Ltd", "Bad vendor", "", 10, "USD", "", "", "", "", ""],
               ["Mars", "Lagos Hardware Co", "Bad company", "", 10, "USD", "Teleported", "tomorrow", "", 50, ""]])
    p = client.post("/api/import/shipments/preview", headers=sales, json={"file": f}).json()
    assert p["summary"]["create"] == 1 and p["summary"]["errors"] == 2
    assert any("Nobody Ltd" in e for e in p["rows"][1]["errors"]) and len(p["rows"][2]["errors"]) >= 3
    assert client.post("/api/import/shipments/commit", headers=sales, json={"file": f}).json()["created"] == 1
    s = next(x for x in client.get("/api/shipments", headers=sales).json() if x["description"] == "Hinges batch")
    assert s["status"] == "in_transit" and s["eta"] == "2026-11-20" and s["vendorPaid"] == 2000 and s["ref"].startswith("IMP-")

    # price lists (admin): new glass thickness in USD, updated ACP price
    cat0 = client.get("/api/quotes/catalog", headers=admin).json()["values"]
    try:
        g = _xlsx([["Glass type", "Thickness", "Price per m²", "Currency"], ["Tempered / Toughened", "15mm", 40, "USD"],
                   ["Tempered / Toughened", "8mm", 19000, ""]])
        assert client.post("/api/import/glass/commit", headers=admin, json={"file": g}).json() == {"created": 1, "updated": 1, "skipped": 0}
        cat = client.get("/api/quotes/catalog", headers=admin).json()["values"]
        assert "15mm" in cat["gc"]["thicknesses"]["tempered"] and cat["gc"]["basePrices"]["tempered__8mm"] == 19000
        assert cat["currencies"]["gc.basePrices.tempered__15mm"] == "USD"
        a = _xlsx([["Panel", "Price"], ["4mm Composite Standard", 18500]])
        assert client.post("/api/import/acp/commit", headers=admin, json={"file": a}).json()["updated"] == 1
    finally:
        for k in ("gc", "acpPanels", "currencies"):
            client.put(f"/api/quotes/catalog/{k}", headers=admin, json={"value": cat0[k] if k != "currencies" else {}})

    bad = {"name": "x.pdf", "data": _b.b64encode(b"%PDF").decode()}
    assert client.post("/api/import/vendors/preview", headers=sales, json={"file": bad}).status_code == 400
    assert "Missing column" in client.post("/api/import/vendors/preview", headers=sales, json={"file": _xlsx([["Country"], ["NG"]])}).json()["detail"]
