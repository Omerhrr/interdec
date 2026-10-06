"""Tests for legacy-facade migration + Excel export.

Run:  cd /home/z/my-project/interdec-platform/backend && python3 -m pytest tests/test_quotes_export_migrate.py -v
"""
import io
import os
import sys
import time
import uuid

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app import models, pricing
from app.database import Base
from app.excel_export import build_quote_workbook, build_quotes_list_workbook
from app.routers import quotes as quotes_router

ENGINE = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
TestSession = sessionmaker(autocommit=False, autoflush=False, bind=ENGINE)

DAY = 86400000
NOW = int(time.time() * 1000)


def _db():
    Base.metadata.create_all(bind=ENGINE)
    db = TestSession()
    yield db
    db.close()


def _user(admin=True):
    return models.User(id="u_test", name="Tester", email=f"t{uuid.uuid4().hex[:6]}@x.io",
                       password_hash="x", platform_role="admin" if admin else "user", apps={}, active=True)


# Facade-shaped payload mirroring the localStorage dataset (idf8:projects)
FACADE_GOLDEN_ITEMS = [{
    "id": "li1", "product": "et_window", "subTypeId": "etw_single",
    "width": 1200, "height": 1500, "qty": 10, "position": "North Elevation",
    "glassTypeId": "tempered", "glassThickness": "8mm", "glassColour": "clear",
    "hasMosq": False, "hasSubframe": False, "elements": [],
}]

FACADE_PAYLOAD = {
    "projects": [{
        "id": "fp1", "name": "Lekki Twin Towers", "siteAddress": "Lekki Phase 1, Lagos",
        "salesPersonId": "sp1",
        "client": {"name": "Towers Development Co.", "phone": "+234 801 111 2222",
                   "email": "pm@towers.ng", "compEmail": "info@towers.ng", "compPhone": "+234 801 333 4444"},
        "quotes": [{
            "id": "fq1", "number": "IDF-2026-090", "description": "Phase 1 windows",
            "createdAt": "2026-02-01",
            "revisions": [{
                "id": "fr1", "rev": 1, "status": "won", "notes": "Client approved",
                "markupOverride": 100, "lineItems": FACADE_GOLDEN_ITEMS,
                "outsideLagos": True, "logisticsLocation": "Ibadan", "logisticsCost": 150000,
                "createdByName": "Chidi Okafor",
                "createdAt": "2026-02-01", "updatedAt": "2026-02-03", "validUntil": "2026-02-15",
            }],
        }],
    }],
    "settings": pricing.SETTINGS_DEF,
    "rates": pricing.RATES_DEF,
    "salesPersons": [{"id": "sp1", "name": "Chidi Okafor"}],
}


def test_migrate_imports_and_preserves_figures():
    db = next(_db())
    res = quotes_router.migrate_legacy(FACADE_PAYLOAD, _user(), db)
    assert res["imported"] == 1 and res["skipped"] == 0
    assert res["importedNumbers"] == ["IDF-2026-090"]

    q = db.query(models.Quote).filter(models.Quote.number == "IDF-2026-090").one()
    assert q.project_name == "Lekki Twin Towers"
    assert q.client_name == "Towers Development Co."
    assert q.client_phone == "+234 801 111 2222"
    assert q.sales_person == "Chidi Okafor"
    assert "Site: Lekki Phase 1, Lagos" in q.notes
    assert "Phase 1 windows" in q.notes

    revs = db.query(models.QuoteRevision).filter(models.QuoteRevision.quote_id == q.id).all()
    assert len(revs) == 1
    r = revs[0]
    assert r.rev == 1 and r.status == "won"
    # golden case figures preserved exactly (engine port parity)
    assert r.cost == 1_666_000
    assert r.mk_amt == 1_666_000
    assert r.sub == 3_332_000
    assert r.vat_amt == 249_900
    # logistics rides on top of VAT
    assert r.logistics_cost == 150_000
    assert r.logistics_location == "Ibadan"
    assert r.total == 3_581_900 + 150_000
    # validity derived from validUntil - createdAt = 14 days
    assert r.validity_days == 14


def test_migrate_is_idempotent_by_number():
    db = next(_db())
    quotes_router.migrate_legacy(FACADE_PAYLOAD, _user(), db)
    res2 = quotes_router.migrate_legacy(FACADE_PAYLOAD, _user(), db)
    assert res2["imported"] == 0
    assert res2["skippedNumbers"] == ["IDF-2026-090"]
    assert db.query(models.Quote).filter(models.Quote.number == "IDF-2026-090").count() == 1


def test_migrate_recomputes_with_facade_rate_overrides():
    """If the legacy app had locally edited rates, totals follow those rates."""
    import copy
    payload = copy.deepcopy(FACADE_PAYLOAD)
    payload["projects"][0]["quotes"][0]["number"] = "IDF-2026-091"
    payload["projects"][0]["quotes"][0]["revisions"][0]["outsideLagos"] = False
    payload["rates"] = copy.deepcopy(pricing.RATES_DEF)
    payload["rates"]["brandRates"]["EUROTEC"] = 3800  # legacy-local override
    db = next(_db())
    res = quotes_router.migrate_legacy(payload, _user(), db)
    assert res["imported"] == 1
    r = db.query(models.QuoteRevision).filter(models.QuoteRevision.quote_id == (
        db.query(models.Quote).filter(models.Quote.number == "IDF-2026-091").one().id)).one()
    # et_window brand IS EUROTEC, so the legacy-local 3800 override raises profile cost
    # by 198 kg x N400 = N79,200 -> cost 1,745,200; total = 1,745,200 x 2 x 1.075
    assert r.cost == 1_745_200
    assert r.logistics_cost == 0 and r.total == 3_752_180


def test_export_quote_workbook_golden():
    rev = {
        "id": "r1", "rev": 1, "status": "draft", "markupPct": 100, "vatPct": 7.5,
        "validityDays": 14, "paymentTerms": pricing.SETTINGS_DEF["paymentTerms"],
        "logisticsCost": 0, "logisticsLocation": "",
        "created": NOW, "cost": 1_666_000, "mkAmt": 1_666_000, "sub": 3_332_000,
        "vatAmt": 249_900, "total": 3_581_900,
        "items": pricing.calc_totals([FACADE_GOLDEN_ITEMS[0]], 100, 7.5)["items"],
    }
    quote = {"number": "IDF-2026-090", "clientName": "Towers Development Co.",
             "projectName": "Lekki Twin Towers", "salesPerson": "Chidi Okafor"}
    data = build_quote_workbook(quote, rev, pricing.DEFAULT_CATALOG)
    assert data[:2] == b"PK"  # valid xlsx (zip)

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Work Breakdown", "BOM Details"]
    ws = wb["Work Breakdown"]
    texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert any("GRAND TOTAL" in t for t in texts)
    grand = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, (int, float))]
    assert 3_581_900 in grand
    # unit price = line selling / qty = 3,332,000 / 10 = 333,200 (legacy double-count fixed)
    assert 333_200 in grand


def test_export_list_workbook_totals():
    quotes = [
        {"number": "IDF-2026-001", "clientName": "A", "projectName": "", "salesPerson": "S",
         "status": "draft", "total": 3_581_900, "updated": NOW,
         "latest": {"rev": 0, "markupPct": 100, "vatPct": 7.5, "logisticsCost": 0},
         "revisions": [{"rev": 0, "status": "draft", "cost": 1, "mkAmt": 1, "sub": 1, "vatAmt": 1,
                        "logisticsCost": 0, "total": 3_581_900, "items": [1, 2, 3]}]},
        {"number": "IDF-2026-002", "clientName": "B", "projectName": "", "salesPerson": "S",
         "status": "won", "total": 1_000_000, "updated": NOW,
         "latest": {"rev": 2, "markupPct": 90, "vatPct": 7.5, "logisticsCost": 50_000},
         "revisions": [{"rev": 2, "status": "won", "cost": 1, "mkAmt": 1, "sub": 1, "vatAmt": 1,
                        "logisticsCost": 50_000, "total": 1_000_000, "items": [1]}]},
    ]
    data = build_quotes_list_workbook(quotes)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Quotes", "Revisions"]
    ws = wb["Quotes"]
    nums = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, (int, float))]
    assert 4_581_900 in nums  # totals row
    assert any(c.value == "IDF-2026-001" for row in ws.iter_rows() for c in row)


if __name__ == "__main__":
    import traceback
    fails = 0
    for name, fn in sorted([(k, v) for k, v in globals().items() if k.startswith("test_")]):
        try:
            fn()
            print(f"\u2713 {name}")
        except Exception:
            fails += 1
            print(f"\u2717 {name}")
            traceback.print_exc()
    sys.exit(1 if fails else 0)
