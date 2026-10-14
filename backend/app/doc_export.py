"""Generic business documents (ImportFlow lists, records and reports) as PDF
(WeasyPrint) or Excel (openpyxl) from one description.

A document is plain data:
    {
      "title": "Shipments", "subtitle": "12 shipments · filtered by In Transit",
      "landscape": True,                       # wide tables prefer landscape
      "summary": [{"label": "In transit", "value": "4", "sub": "$ 120,000", "color": "#f59e0b"}],
      "sections": [{
          "heading": "Mentor House", "note": "$ 40,000 · 3 shipments", "color": "#3b82f6",
          "kv": [["Status", "In Transit"], ...],          # label/value pairs (record pages)
          "columns": [{"label": "Ref"}, {"label": "Value", "align": "right", "num": True}],
          "rows": [["SP-001", 1200.5], ...],
          "total": ["Total", "", 1200.5],                   # optional footer row
      }],
      "notes": ["Values converted to USD at ..."],
    }
Numbers stay numbers in Excel (so they can be summed); the PDF formats them.
The look comes from the ImportFlow "Document Design" settings.
"""
import base64
import io
from datetime import datetime

from jinja2 import Environment, BaseLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .pdf_export import FONTS, HEX, LAYOUTS, PAPERS, PdfUnavailable, _data_only_fetcher, validate_design as _validate_base

DEFAULT_DESIGN = {
    "companyName": "Interdec",
    "tagline": "ImportFlow",
    "address": "",
    "phone": "",
    "email": "",
    "website": "",
    "logo": "",
    "logoHeight": 36,
    "brandColor": "#1E3A5F",
    "accentColor": "#3B82F6",
    "font": "sans",
    "layout": "classic",
    "paper": "A4",
    "orientation": "auto",      # auto (wide tables go landscape) | portrait | landscape
    "footerText": "Confidential · for internal use",
    "show": {
        "summary": True,         # coloured stat cards at the top
        "pageNumbers": True,
        "generated": True,       # "Generated 13 Oct 2026 by …"
        "stripes": True,         # zebra rows
        "notes": True,           # exchange-rate notes etc.
    },
}
ORIENTATIONS = ("auto", "portrait", "landscape")


def merged_design(saved):
    d = {**DEFAULT_DESIGN, **{k: v for k, v in (saved or {}).items() if k != "show"}}
    d["show"] = {**DEFAULT_DESIGN["show"], **((saved or {}).get("show") or {})}
    return d


def validate_design(value) -> str | None:
    msg = _validate_base(value)
    if msg:
        return msg
    if value.get("orientation") and value["orientation"] not in ORIENTATIONS:
        return f"orientation must be one of {list(ORIENTATIONS)}"
    return None


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def fmt_cell(v, col):
    if v is None or v == "":
        return ""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        dec = col.get("decimals", 0)
        s = f"{v:,.{dec}f}"
        return f"{col.get('prefix', '')}{s}{col.get('suffix', '')}"
    return str(v)


TEMPLATE = r"""<!doctype html><html><head><meta charset="utf-8"><title>{{ doc.title }}</title>
<style>
@page {
  size: {{ d.paper }} {{ orient }}; margin: 14mm 12mm 16mm;
  {% if footer %}@bottom-left { content: "{{ footer|safe }}"; font-family: {{ font|safe }}; font-size: 7pt; color: #7a8594; }{% endif %}
  {% if d.show.pageNumbers %}@bottom-right { content: "Page " counter(page) " of " counter(pages); font-family: {{ font|safe }}; font-size: 7pt; color: #7a8594; }{% endif %}
}
* { box-sizing: border-box; }
body { font-family: {{ font|safe }}; font-size: 8.6pt; color: #1f2937; margin: 0; line-height: 1.38; }
.head { display: flex; justify-content: space-between; align-items: flex-start; gap: 16px; padding-bottom: 10px; margin-bottom: 12px; }
.head > div:first-child { flex: 1; min-width: 0; }
.classic .head { border-bottom: 2.2pt solid {{ d.brandColor }}; }
.minimal .head { border-bottom: 0.6pt solid #d6dce4; }
.banner .head { background: {{ d.brandColor }}; color: #fff; margin: -14mm -12mm 14px; padding: 11mm 12mm 7mm; }
.banner .co, .banner .title { color: #fff; }
.banner .muted { color: rgba(255,255,255,.8); }
.logo { height: {{ d.logoHeight }}px; max-width: 200px; object-fit: contain; display: block; margin-bottom: 4px; }
.co { font-size: 12.5pt; font-weight: 700; color: {{ d.brandColor }}; }
.muted { color: #64748b; }
.small { font-size: 7.6pt; }
.right { text-align: right; flex: 0 0 auto; max-width: 55%; }
.title { font-size: 16pt; font-weight: 800; color: {{ d.brandColor }}; letter-spacing: -.2px; }
.minimal .title { font-weight: 400; letter-spacing: 1px; }
.cards { display: flex; gap: 8px; margin: 0 0 12px; }
.card { flex: 1; border: 0.6pt solid #e2e8f0; border-top: 2.4pt solid {{ d.accentColor }}; border-radius: 4px; padding: 7px 9px; }
.card .v { font-size: 13pt; font-weight: 800; }
.card .l { font-size: 7.4pt; font-weight: 700; color: #475569; text-transform: uppercase; letter-spacing: .5px; }
.card .s { font-size: 7.6pt; color: #64748b; }
.sec { margin-top: 12px; }
.sec h2 { font-size: 10.5pt; margin: 0 0 2px; color: {{ d.brandColor }}; page-break-after: avoid; }
.sec .note { font-size: 8pt; color: #475569; margin-bottom: 5px; page-break-after: avoid; }
.kv { display: grid; grid-template-columns: repeat(3, 1fr); gap: 6px 16px; margin-bottom: 8px; }
.kv div b { display: block; font-size: 7pt; text-transform: uppercase; letter-spacing: .5px; color: #64748b; font-weight: 700; }
table { width: 100%; border-collapse: collapse; }
thead { display: table-header-group; }
th { background: {{ d.brandColor }}; color: #fff; font-size: 7pt; text-transform: uppercase; letter-spacing: .4px; padding: 5px 6px; text-align: left; font-weight: 700; }
.minimal th { background: none; color: #475569; border-bottom: 1.1pt solid {{ d.brandColor }}; }
td { padding: 5px 6px; border-bottom: 0.5pt solid #e2e8f0; vertical-align: top; }
tr { page-break-inside: avoid; }
{% if d.show.stripes %}tbody tr:nth-child(even) td { background: #f8fafc; }{% endif %}
tfoot td:first-child { white-space: nowrap; }
tfoot td { font-weight: 800; border-top: 1pt solid {{ d.brandColor }}; border-bottom: none; }
.r { text-align: right; white-space: nowrap; }
.c { text-align: center; }
.empty { color: #94a3b8; font-style: italic; padding: 8px 0; }
.notes { margin-top: 14px; font-size: 7.6pt; color: #64748b; }
</style></head>
<body class="{{ d.layout }}">
<div class="head">
  <div>
    {% if d.logo %}<img class="logo" src="{{ d.logo }}" alt="">{% endif %}
    {% if not d.logo or d.layout != 'minimal' %}<div class="co">{{ d.companyName }}</div>{% endif %}
    {% if d.tagline %}<div class="muted small">{{ d.tagline }}</div>{% endif %}
    {% set contact = [d.address, d.phone, d.email, d.website]|select|list %}
    {% if contact %}<div class="muted small">{{ contact|join("  ·  ") }}</div>{% endif %}
  </div>
  <div class="right">
    <div class="title">{{ doc.title }}</div>
    {% if doc.subtitle %}<div class="small">{{ doc.subtitle }}</div>{% endif %}
    {% if d.show.generated %}<div class="muted small">Generated {{ now.strftime("%d %b %Y, %H:%M") }}{% if by %} by {{ by }}{% endif %}</div>{% endif %}
  </div>
</div>

{% if d.show.summary and doc.summary %}
<div class="cards">
  {% for c in doc.summary %}
  <div class="card" {% if c.color %}style="border-top-color: {{ c.color }}"{% endif %}>
    <div class="v" {% if c.color %}style="color: {{ c.color }}"{% endif %}>{{ c.value }}</div>
    <div class="l">{{ c.label }}</div>
    {% if c.sub %}<div class="s">{{ c.sub }}</div>{% endif %}
  </div>
  {% endfor %}
</div>
{% endif %}

{% for s in doc.sections %}
<div class="sec">
  {% if s.heading %}<h2 {% if s.color %}style="color: {{ s.color }}"{% endif %}>{{ s.heading }}</h2>{% endif %}
  {% if s.note %}<div class="note">{{ s.note }}</div>{% endif %}
  {% if s.kv %}<div class="kv">{% for k, v in s.kv %}<div><b>{{ k }}</b>{{ v if v not in (None, "") else "—" }}</div>{% endfor %}</div>{% endif %}
  {% if s.columns %}
    {% if s.rows %}
    <table>
      <thead><tr>{% for c in s.columns %}<th class="{{ 'r' if c.align == 'right' else 'c' if c.align == 'center' else '' }}">{{ c.label }}</th>{% endfor %}</tr></thead>
      <tbody>{% for r in s.rows %}<tr>{% for c in s.columns %}<td class="{{ 'r' if c.align == 'right' else 'c' if c.align == 'center' else '' }}">{{ fmt(r[loop.index0], c) }}</td>{% endfor %}</tr>{% endfor %}</tbody>
      {% if s.total %}<tfoot><tr>{% for c in s.columns %}<td class="{{ 'r' if c.align == 'right' else '' }}">{{ fmt(s.total[loop.index0], c) }}</td>{% endfor %}</tr></tfoot>{% endif %}
    </table>
    {% else %}<div class="empty">{{ s.empty or "Nothing to show." }}</div>{% endif %}
  {% endif %}
</div>
{% endfor %}

{% if d.show.notes and doc.notes %}<div class="notes">{% for n in doc.notes %}<div>{{ n }}</div>{% endfor %}</div>{% endif %}
</body></html>"""

_env = Environment(loader=BaseLoader(), autoescape=select_autoescape(default=True, default_for_string=True))
_tpl = _env.from_string(TEMPLATE)


def _orientation(doc, d):
    if d["orientation"] in ("portrait", "landscape"):
        return d["orientation"]
    return "landscape" if doc.get("landscape") else "portrait"


def render_html(doc: dict, design: dict, by: str = "") -> str:
    import re
    d = merged_design(design)
    footer = re.sub(r'[\\"<>\n\r]', " ", d.get("footerText") or "")[:200]
    return _tpl.render(doc=doc, d=d, by=by, now=datetime.now(), fmt=fmt_cell, footer=footer,
                       font=FONTS.get(d["font"], FONTS["sans"]), orient=_orientation(doc, d))


def render_pdf(doc: dict, design: dict, by: str = "") -> bytes:
    try:
        from weasyprint import HTML
    except Exception as e:  # pragma: no cover - depends on the OS
        raise PdfUnavailable(
            "PDF export needs WeasyPrint. Install it with `pip install weasyprint` "
            "(and on Linux/WSL: `sudo apt install libpango-1.0-0 libpangoft2-1.0-0`)."
        ) from e
    return HTML(string=render_html(doc, design, by), url_fetcher=_data_only_fetcher()).write_pdf()


# ---------------------------------------------------------------------------
# Excel: one sheet per document, sections stacked, numbers kept numeric
# ---------------------------------------------------------------------------

def _hex(c, fallback="1E3A5F"):
    c = (c or "").lstrip("#")
    return c.upper() if len(c) == 6 else fallback


def _logo_image(d):
    logo = d.get("logo") or ""
    if not logo.startswith("data:image/") or "svg" in logo[:30]:
        return None  # Excel can't embed SVG
    try:
        raw = base64.b64decode(logo.split(",", 1)[1])
        img = XLImage(io.BytesIO(raw))
        h = 40
        img.width, img.height = img.width * h / max(img.height, 1), h
        return img
    except Exception:
        return None


def _sheet_name(s):
    bad = '[]:*?/\\'
    return "".join("-" if ch in bad else ch for ch in s)[:31] or "Sheet"


def render_xlsx(doc: dict, design: dict, by: str = "") -> bytes:
    d = merged_design(design)
    brand, accent = _hex(d["brandColor"]), _hex(d["accentColor"], "3B82F6")
    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_name(doc.get("sheet") or doc["title"])
    thin = Side(style="thin", color="E2E8F0")
    width = max([len(s.get("columns") or []) for s in doc.get("sections") or []] + [4])

    row = 1
    img = _logo_image(d)
    if img:
        ws.add_image(img, "A1")
        ws.row_dimensions[1].height = 32
        row = 3
    ws.cell(row, 1, d["companyName"]).font = Font(bold=True, size=12, color=brand)
    row += 1
    ws.cell(row, 1, doc["title"]).font = Font(bold=True, size=15, color="0F172A")
    row += 1
    sub = " · ".join(x for x in [doc.get("subtitle"), f"Generated {datetime.now():%d %b %Y %H:%M}" + (f" by {by}" if by else "")] if x)
    ws.cell(row, 1, sub).font = Font(size=9, color="64748B")
    row += 2

    if d["show"]["summary"] and doc.get("summary"):
        for i, c in enumerate(doc["summary"]):
            col = 1 + i * 2
            ws.cell(row, col, c["label"]).font = Font(bold=True, size=8, color="475569")
            v = ws.cell(row + 1, col, c["value"])
            v.font = Font(bold=True, size=13, color=_hex(c.get("color"), accent))
            if c.get("sub"):
                ws.cell(row + 2, col, c["sub"]).font = Font(size=9, color="64748B")
        row += 4

    for s in doc.get("sections") or []:
        if s.get("heading"):
            ws.cell(row, 1, s["heading"]).font = Font(bold=True, size=11, color=_hex(s.get("color"), brand))
            row += 1
        if s.get("note"):
            ws.cell(row, 1, s["note"]).font = Font(size=9, color="475569")
            row += 1
        for k, v in s.get("kv") or []:
            ws.cell(row, 1, k).font = Font(bold=True, size=9, color="64748B")
            ws.cell(row, 2, "—" if v in (None, "") else v)
            row += 1
        cols = s.get("columns") or []
        if cols:
            for ci, c in enumerate(cols, start=1):
                h = ws.cell(row, ci, c["label"])
                h.font = Font(bold=True, color="FFFFFF", size=9)
                h.fill = PatternFill("solid", fgColor=brand)
                h.alignment = Alignment(horizontal="right" if c.get("align") == "right" else "left", vertical="center")
            ws.row_dimensions[row].height = 18
            row += 1
            for ri, r in enumerate(s.get("rows") or []):
                for ci, c in enumerate(cols, start=1):
                    v = r[ci - 1] if ci - 1 < len(r) else None
                    cell = ws.cell(row, ci, v)
                    cell.border = Border(bottom=thin)
                    if c.get("num") and isinstance(v, (int, float)):
                        cell.number_format = "#,##0" + ("." + "0" * c["decimals"] if c.get("decimals") else "")
                        cell.alignment = Alignment(horizontal="right")
                    if d["show"]["stripes"] and ri % 2:
                        cell.fill = PatternFill("solid", fgColor="F8FAFC")
                row += 1
            if not s.get("rows"):
                ws.cell(row, 1, s.get("empty") or "Nothing to show.").font = Font(italic=True, color="94A3B8")
                row += 1
            if s.get("total"):
                for ci, c in enumerate(cols, start=1):
                    v = s["total"][ci - 1] if ci - 1 < len(s["total"]) else None
                    cell = ws.cell(row, ci, v)
                    cell.font = Font(bold=True)
                    cell.border = Border(top=Side(style="medium", color=brand))
                    if c.get("num") and isinstance(v, (int, float)):
                        cell.number_format = "#,##0" + ("." + "0" * c["decimals"] if c.get("decimals") else "")
                row += 1
        row += 1

    if d["show"]["notes"]:
        for n in doc.get("notes") or []:
            ws.cell(row, 1, n).font = Font(size=8, italic=True, color="64748B")
            row += 1

    # column widths from content
    for ci in range(1, width + 1):
        longest = 8
        for r in ws.iter_rows(min_col=ci, max_col=ci, min_row=8):
            v = r[0].value
            if v is not None:
                longest = max(longest, min(len(f"{v:,.2f}") if isinstance(v, float) else len(str(v)), 48))
        ws.column_dimensions[get_column_letter(ci)].width = longest + 2
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = _orientation(doc, d)
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
