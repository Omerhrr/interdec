"""Bring existing records in from Excel (.xlsx) or CSV.

Each kind has a template (downloadable), flexible header matching (case and
spacing don't matter; common alternative names are accepted), a preview that
shows per-row errors and whether the row will create or update, and a commit
that applies only the valid rows. Rows match existing records by name (or
shipment ref), so re-importing a corrected file updates instead of duplicating.
"""
import base64
import csv
import io
import re
import time
import uuid
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from . import models

MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024
COMPANIES = {"facade": "facade", "davinci": "davinci", "doortec": "doortec"}
STATUSES = {"order placed": "order_placed", "under production": "under_production", "in transit": "in_transit",
            "completed": "completed", "order_placed": "order_placed", "under_production": "under_production",
            "in_transit": "in_transit"}

# kind -> (title, columns [(key, header, required, aliases, example)], who can import)
CONTACT_COLS = [
    ("name", "Name", True, ["company", "company name", "vendor", "supplier"], ""),
    ("contact", "Contact person", False, ["contact", "contact name"], ""),
    ("email", "Email", False, ["e-mail", "email address"], ""),
    ("phone", "Phone", False, ["telephone", "tel", "mobile"], ""),
    ("country", "Country", False, [], ""),
    ("category", "Category", False, ["type"], ""),
    ("notes", "Notes", False, ["note", "comments"], ""),
]
KINDS = {
    "vendors": {"title": "Vendors", "who": "importflow",
                "cols": [(k, h, r, a, e) for k, h, r, a, e in CONTACT_COLS] + [("address", "Address", False, [], "")],
                "example": ["Vetro Italia Glass", "Marco Rossi", "marco@vetro.it", "+39 02 555 0100", "Italy", "Glass & Glazing", "Tempered glass, 6 weeks", "Via Roma 1, Milan"]},
    "shippers": {"title": "Shippers", "who": "importflow", "cols": CONTACT_COLS,
                 "example": ["Global Freight Logistics", "Ahmed Hassan", "ahmed@globalfreight.ae", "+971 4 567 8901", "UAE", "Ocean Freight", ""]},
    "agents": {"title": "Clearing agents", "who": "importflow", "cols": CONTACT_COLS,
               "example": ["Apapa Clearing & Forwarding", "Chinedu Okeke", "ops@apapaclearing.ng", "+234 803 555 0101", "Nigeria", "Sea port", ""]},
    "projects": {"title": "Projects", "who": "importflow",
                 "cols": [("name", "Project name", True, ["project", "name"], ""), ("client", "Client", False, ["customer"], ""),
                          ("status", "Status", False, [], ""), ("description", "Description", False, ["notes"], "")],
                 "example": ["Lekki Towers Phase 2", "Lekki Towers Ltd", "Active", "Curtain wall and windows"]},
    "shipments": {"title": "Shipments", "who": "importflow",
                  "cols": [("ref", "Ref", False, ["reference", "shipment ref", "shipment"], ""),
                           ("company", "Company", True, [], ""), ("vendor", "Vendor", True, ["supplier"], ""),
                           ("description", "Description", True, ["goods", "item"], ""), ("category", "Category", False, [], ""),
                           ("project", "Project", False, [], ""), ("value", "Value", False, ["amount", "order value"], ""),
                           ("currency", "Currency", False, ["ccy"], ""), ("status", "Status", False, ["stage"], ""),
                           ("eta", "ETA", False, ["arrival", "eta date"], ""), ("shipper", "Shipper", False, ["forwarder"], ""),
                           ("method", "Shipping method", False, ["method", "ship method"], ""),
                           ("terms", "Payment terms", False, ["terms"], ""), ("production", "Production days", False, ["production time", "lead time"], ""),
                           ("paid", "Paid to vendor", False, ["amount paid", "paid"], "")],
                  "example": ["", "Facade", "Vetro Italia Glass", "Glass panels batch 3", "Glass & Glazing", "Lekki Towers Phase 2",
                              "48500", "USD", "In Transit", "2026-11-20", "Global Freight Logistics", "FCL", "After Shipment", "45", "14550"]},
    "glass": {"title": "Glass prices", "who": "admin",
              "cols": [("type", "Glass type", True, ["glass"], ""), ("thickness", "Thickness", True, ["size"], ""),
                       ("price", "Price per m2", True, ["price", "rate", "price per m²"], ""), ("currency", "Currency", False, [], ""),
                       ("dgu", "Double glazed", False, ["dgu"], "")],
              "example": ["Tempered / Toughened", "8mm", "18500", "NGN", "No"]},
    "acp": {"title": "ACP panels", "who": "admin",
            "cols": [("name", "Panel", True, ["name", "panel name"], ""), ("price", "Price per m2", True, ["price", "rate"], ""),
                     ("currency", "Currency", False, [], "")],
            "example": ["4mm Composite Standard", "17000", "NGN"]},
    "hardware": {"title": "Hardware prices", "who": "admin",
                 "cols": [("name", "Item", True, ["name", "hardware"], ""), ("unit", "Unit", False, [], ""),
                          ("price", "Price", True, ["unit price", "rate"], ""), ("currency", "Currency", False, [], "")],
                 "example": ["Window Handle", "piece", "5000", "NGN"]},
}


def norm(s) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower().replace("²", "2")).strip()


def template(kind: str) -> bytes:
    spec = KINDS[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = spec["title"][:31]
    for i, (_, h, req, _, _) in enumerate(spec["cols"], start=1):
        c = ws.cell(1, i, h + (" *" if req else ""))
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A5F")
        ws.column_dimensions[c.column_letter].width = max(14, len(h) + 4)
    for i, v in enumerate(spec["example"], start=1):
        ws.cell(2, i, v).font = Font(italic=True, color="64748B")
    help_ = wb.create_sheet("How to fill")
    lines = [f"Import {spec['title'].lower()} into Interdec", "",
             "• Keep the header row. Columns marked * are required; the order doesn't matter.",
             "• Row 2 is an example: replace or delete it.",
             "• A row whose name (or shipment Ref) already exists updates that record instead of adding a copy."]
    if kind == "shipments":
        lines += ["• Company: Facade, Davinci or Doortec.  Status: Order Placed, Under Production, In Transit or Completed.",
                  "• Vendor, Project and Shipper must already exist (import them first) — names are matched exactly, ignoring case.",
                  "• ETA: a date (2026-11-20 or 20/11/2026).  Paid to vendor: amount already paid, recorded as one payment.",
                  "• Leave Ref empty for new shipments; a reference is assigned."]
    if kind in ("glass", "acp", "hardware"):
        lines += ["• Currency is optional (default NGN). Prices replace the current ones; new names are added."]
    if kind == "glass":
        lines += ["• A new glass type or thickness is created when the name doesn't exist yet."]
    for i, l in enumerate(lines, start=1):
        help_.cell(i, 1, l).font = Font(bold=(i == 1), size=12 if i == 1 else 11)
    help_.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_rows(file: dict) -> list[list]:
    if not isinstance(file, dict) or not file.get("data"):
        raise ValueError("Upload a file")
    name = (file.get("name") or "").lower()
    try:
        raw = base64.b64decode(file["data"].split(",")[-1], validate=False)
    except Exception:
        raise ValueError("The file couldn't be read")
    if len(raw) > MAX_BYTES:
        raise ValueError("File is larger than 5 MB")
    if name.endswith(".csv") or name.endswith(".txt"):
        text = raw.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t") if text.strip() else csv.excel
        rows = [r for r in csv.reader(io.StringIO(text), dialect)]
    elif name.endswith(".xlsx") or name.endswith(".xlsm"):
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:
            raise ValueError("That isn't a readable Excel (.xlsx) file")
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        raise ValueError("Use an Excel (.xlsx) or CSV file")
    rows = [r for r in rows if any(c not in (None, "") and str(c).strip() for c in r)]
    if len(rows) > MAX_ROWS + 1:
        raise ValueError(f"Up to {MAX_ROWS} rows per import")
    return rows


def map_headers(kind, header) -> tuple[dict, list[str]]:
    spec = KINDS[kind]
    idx, missing = {}, []
    hn = [norm(str(h).replace("*", "")) for h in header]
    for key, h, req, aliases, _ in spec["cols"]:
        names = {norm(h)} | {norm(a) for a in aliases}
        pos = next((i for i, x in enumerate(hn) if x in names), None)
        if pos is not None:
            idx[key] = pos
        elif req:
            missing.append(h)
    return idx, missing


def _num(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.\-]", "", str(v))
    try:
        return float(s)
    except ValueError:
        raise ValueError(f"'{v}' isn't a number")


def _date(v):
    if v in (None, ""):
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d.%m.%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(s, f).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"'{v}' isn't a date (use 2026-11-20 or 20/11/2026)")


def _s(v, n=300):
    return ("" if v is None else str(v)).strip()[:n]


# ---------------------------------------------------------------------------
# validate + apply
# ---------------------------------------------------------------------------

class Ctx:
    def __init__(self, db, catalog=None, fx_codes=None):
        self.db = db
        self.catalog = catalog or {}
        self.fx = {c.upper() for c in (fx_codes or ["NGN", "USD"])} | {"NGN", "USD"}
        self.by_name = {
            "vendors": {norm(x.name): x for x in db.query(models.Vendor).all()},
            "shippers": {norm(x.name): x for x in db.query(models.Shipper).all()},
            "agents": {norm(x.name): x for x in db.query(models.ClearingAgent).all()},
            "projects": {norm(x.name): x for x in db.query(models.Project).all()},
        }
        self.ships = {norm(s.ref): s for s in db.query(models.Shipment).all()}


def _cur(v, ctx):
    c = _s(v, 3).upper() or "NGN"
    if c not in ctx.fx:
        raise ValueError(f"No exchange rate for currency {c}")
    return c


def check_row(kind, rec, ctx, seen) -> tuple[str, list[str], dict]:
    """-> (action create|update, errors, cleaned values)"""
    errs, clean = [], {}

    def need(key, label):
        if not _s(rec.get(key)):
            errs.append(f"{label} is required")

    try:
        if kind in ("vendors", "shippers", "agents", "projects"):
            need("name", "Name")
            clean = {k: _s(rec.get(k), 2000 if k in ("notes", "description") else 200) for k in rec}
            if kind == "projects":
                st = _s(rec.get("status")).title() or "Active"
                if st == "On hold":
                    st = "On Hold"
                if st not in ("Active", "On Hold", "Completed"):
                    errs.append("Status must be Active, On Hold or Completed")
                clean["status"] = st
            key = norm(clean.get("name"))
            action = "update" if key in ctx.by_name[kind] else "create"
        elif kind == "shipments":
            for k, l in (("company", "Company"), ("vendor", "Vendor"), ("description", "Description")):
                need(k, l)
            comp = COMPANIES.get(norm(rec.get("company")))
            if _s(rec.get("company")) and not comp:
                errs.append("Company must be Facade, Davinci or Doortec")
            vendor = ctx.by_name["vendors"].get(norm(rec.get("vendor")))
            if _s(rec.get("vendor")) and not vendor:
                errs.append(f"Vendor '{_s(rec.get('vendor'))}' not found (import vendors first)")
            proj = ctx.by_name["projects"].get(norm(rec.get("project"))) if _s(rec.get("project")) else None
            if _s(rec.get("project")) and not proj and norm(rec.get("project")) not in ("general stock", "none"):
                errs.append(f"Project '{_s(rec.get('project'))}' not found")
            shipper = ctx.by_name["shippers"].get(norm(rec.get("shipper"))) if _s(rec.get("shipper")) else None
            if _s(rec.get("shipper")) and not shipper:
                errs.append(f"Shipper '{_s(rec.get('shipper'))}' not found")
            status = STATUSES.get(norm(rec.get("status")) or "order placed")
            if not status:
                errs.append("Status must be Order Placed, Under Production, In Transit or Completed")
            terms = _s(rec.get("terms")).title()
            if terms and terms not in ("Before Shipment", "After Shipment"):
                errs.append("Payment terms must be Before Shipment or After Shipment")
            value = _num(rec.get("value")) or 0.0
            paid = _num(rec.get("paid")) or 0.0
            if value < 0 or paid < 0:
                errs.append("Value and paid can't be negative")
            if paid > value + 0.005:
                errs.append("Paid to vendor is more than the order value")
            prod = _num(rec.get("production"))
            clean = {"company": comp, "vendor": vendor.id if vendor else None, "project": proj.id if proj else None,
                     "shipper": shipper.id if shipper else None, "status": status, "terms": terms or None,
                     "description": _s(rec.get("description")), "category": _s(rec.get("category"), 80),
                     "value": value, "currency": _cur(rec.get("currency") or "USD", ctx), "eta": _date(rec.get("eta")),
                     "method": _s(rec.get("method"), 40), "production": f"{int(prod)} days" if prod else "", "paid": paid,
                     "ref": _s(rec.get("ref"), 40)}
            action = "update" if clean["ref"] and norm(clean["ref"]) in ctx.ships else "create"
            if clean["ref"] and action == "create" and norm(clean["ref"]) in seen:
                errs.append("Ref appears twice in the file")
            key = norm(clean["ref"]) if clean["ref"] else f"__new{len(seen)}"
        elif kind == "glass":
            need("type", "Glass type"); need("thickness", "Thickness")
            price = _num(rec.get("price"))
            if price is None or price < 0:
                errs.append("Price per m² is required (0 or more)")
            gc = ctx.catalog.get("gc") or {}
            t = next((x for x in gc.get("types", []) if norm(x["label"]) == norm(rec.get("type"))), None)
            th = _s(rec.get("thickness"), 40)
            exists = bool(t and th in (gc.get("thicknesses", {}).get(t["id"]) or []))
            clean = {"type": _s(rec.get("type"), 60), "thickness": th, "price": price, "currency": _cur(rec.get("currency"), ctx),
                     "dgu": norm(rec.get("dgu")) in ("yes", "y", "true", "1", "dgu")}
            action, key = ("update" if exists else "create"), f"{norm(clean['type'])}|{norm(th)}"
        elif kind in ("acp", "hardware"):
            need("name", "Panel" if kind == "acp" else "Item")
            price = _num(rec.get("price"))
            if price is None or price < 0:
                errs.append("Price is required (0 or more)")
            lst = ctx.catalog.get("acpPanels" if kind == "acp" else "hwItems") or []
            clean = {"name": _s(rec.get("name"), 80), "price": price, "currency": _cur(rec.get("currency"), ctx),
                     "unit": _s(rec.get("unit"), 20) or "piece"}
            action = "update" if any(norm(x["name"]) == norm(clean["name"]) for x in lst) else "create"
            key = norm(clean["name"])
        else:
            raise ValueError("Unknown import")
    except ValueError as e:
        errs.append(str(e))
        key, action = f"__err{len(seen)}", "create"
    if key in seen and not key.startswith("__"):
        errs.append("Same name appears twice in the file")
    seen.add(key)
    return action, errs, clean


def analyse(kind, rows, ctx):
    if not rows:
        raise ValueError("The file is empty")
    idx, missing = map_headers(kind, rows[0])
    if missing:
        raise ValueError("Missing column" + ("s" if len(missing) > 1 else "") + ": " + ", ".join(missing) +
                         ". Download the template to see the expected headers.")
    spec = KINDS[kind]
    out, seen = [], set()
    for n, r in enumerate(rows[1:], start=2):
        rec = {k: (r[i] if i < len(r) else None) for k, i in idx.items()}
        # skip the template's untouched example row
        if [_s(rec.get(c[0])) for c in spec["cols"]] == [_s(x) for x in spec["example"]][:len(spec["cols"])]:
            continue
        action, errs, clean = check_row(kind, rec, ctx, seen)
        out.append({"row": n, "values": {c[1]: _s(rec.get(c[0]), 120) for c in spec["cols"] if c[0] in idx},
                    "action": action, "errors": errs, "clean": clean})
    return out


def apply(kind, analysed, ctx, user):
    db, now = ctx.db, int(time.time() * 1000)
    done = {"created": 0, "updated": 0, "skipped": 0}
    good = [r for r in analysed if not r["errors"]]
    done["skipped"] = len(analysed) - len(good)
    if kind in ("vendors", "shippers", "agents", "projects"):
        Model = {"vendors": models.Vendor, "shippers": models.Shipper, "agents": models.ClearingAgent, "projects": models.Project}[kind]
        fields = {c[0] for c in KINDS[kind]["cols"]}
        for r in good:
            c = r["clean"]
            obj = ctx.by_name[kind].get(norm(c["name"]))
            vals = {k: v for k, v in c.items() if k in fields and hasattr(Model, k) and v not in (None, "")}
            if obj:
                for k, v in vals.items():
                    if k != "name":  # matched by name: keep the existing spelling
                        setattr(obj, k, v)
                done["updated"] += 1
            else:
                obj = Model(id=uuid.uuid4().hex[:12], created=now, **{k: v for k, v in vals.items()})
                if kind == "projects" and not obj.status:
                    obj.status = "Active"
                db.add(obj)
                ctx.by_name[kind][norm(c["name"])] = obj
                done["created"] += 1
    elif kind == "shipments":
        from .routers.shipments import _next_ref
        for r in good:
            c = r["clean"]
            s = ctx.ships.get(norm(c["ref"])) if c["ref"] else None
            if not s:
                db.flush()
                s = models.Shipment(id=uuid.uuid4().hex[:12], ref=c["ref"] or _next_ref(db), created=now, dates={"created": now},
                                    goods_confirmed=c["status"] == "completed")
                db.add(s)
                ctx.ships[norm(s.ref)] = s
                done["created"] += 1
            else:
                done["updated"] += 1
            s.company_id, s.vendor_id, s.project_id = c["company"], c["vendor"], c["project"]
            s.description, s.category, s.value, s.currency = c["description"], c["category"], c["value"], c["currency"]
            s.status, s.payment_terms, s.eta, s.ship_method = c["status"], c["terms"], c["eta"], c["method"]
            if c["shipper"]:
                s.shipper_id = c["shipper"]
            if c["production"]:
                s.production_days = c["production"]
            if c["paid"] and not (s.vendor_payments or []):
                s.vendor_payments = [{"id": uuid.uuid4().hex[:10], "date": date.today().isoformat(), "amount": round(c["paid"], 2),
                                      "kind": "Full payment" if c["paid"] >= c["value"] else "Deposit", "ref": "Imported",
                                      "note": "Imported from spreadsheet", "by": user.name, "ts": now}]
            if c["status"] == "completed":
                s.goods_confirmed = True
            s.updated = now
    return done


def apply_prices(kind, analysed, catalog, user):
    """Returns (catalog key, new value, currencies update dict, counts)."""
    import copy
    done = {"created": 0, "updated": 0, "skipped": len([r for r in analysed if r["errors"]])}
    cur = dict(catalog.get("currencies") or {})
    good = [r["clean"] for r in analysed if not r["errors"]]
    if kind == "glass":
        gc = copy.deepcopy(catalog["gc"])
        for c in good:
            t = next((x for x in gc["types"] if norm(x["label"]) == norm(c["type"])), None)
            if not t:
                t = {"id": f"g_{re.sub('[^a-z0-9]+', '_', c['type'].lower())[:20]}_{uuid.uuid4().hex[:4]}", "label": c["type"], "isDGU": c["dgu"], "custom": True}
                gc["types"].append(t)
                gc["thicknesses"][t["id"]] = []
            lst = gc["thicknesses"].setdefault(t["id"], [])
            key = f"{t['id']}__{c['thickness']}"
            if c["thickness"] in lst:
                done["updated"] += 1
            else:
                lst.append(c["thickness"])
                done["created"] += 1
            gc["basePrices"][key] = c["price"]
            path = f"gc.basePrices.{key}"
            if c["currency"] == "NGN":
                cur.pop(path, None)
            else:
                cur[path] = c["currency"]
        return "gc", gc, cur, done
    key = "acpPanels" if kind == "acp" else "hwItems"
    lst = copy.deepcopy(catalog.get(key) or [])
    prefix = "ap" if kind == "acp" else "hw"
    for c in good:
        item = next((x for x in lst if norm(x["name"]) == norm(c["name"])), None)
        if item:
            item["p"] = c["price"]
            if kind == "hardware" and c["unit"]:
                item["unit"] = c["unit"]
            done["updated"] += 1
        else:
            item = {"id": f"{prefix}_{uuid.uuid4().hex[:6]}", "name": c["name"], "p": c["price"], "custom": True}
            if kind == "hardware":
                item["unit"] = c["unit"]
            lst.append(item)
            done["created"] += 1
        path = f"{key}.{item['id']}"
        if c["currency"] == "NGN":
            cur.pop(path, None)
        else:
            cur[path] = c["currency"]
    return key, lst, cur, done
